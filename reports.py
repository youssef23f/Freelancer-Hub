import csv
import os
from datetime import date, timedelta
from core.earnings import get_earnings, get_earnings_stats
from core.time_tracker import get_time_entries, get_time_stats
from core.projects_clients_goals import get_projects, get_clients


def _get_report_dir():
    d = os.path.join(os.path.expanduser("~"), "FreelancerHub_Reports")
    os.makedirs(d, exist_ok=True)
    return d


def export_csv(user_id: int, report_type: str = "monthly") -> str:
    earnings = get_earnings(user_id)
    path = os.path.join(_get_report_dir(), f"earnings_{report_type}_{date.today()}.csv")
    with open(path, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["id","project_name","client_name","amount","date","category","notes"])
        writer.writeheader()
        for e in earnings:
            writer.writerow({k: e.get(k, "") for k in ["id","project_name","client_name","amount","date","category","notes"]})
    return path


def export_excel(user_id: int, report_type: str = "monthly") -> str:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        return export_csv(user_id, report_type)

    wb = openpyxl.Workbook()

    # ── Earnings sheet ──
    ws = wb.active
    ws.title = "Earnings"
    header_fill = PatternFill(start_color="1a2235", end_color="1a2235", fill_type="solid")
    header_font = Font(bold=True, color="60a5fa")
    headers = ["Project", "Client", "Amount ($)", "Date", "Category", "Notes"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row, e in enumerate(get_earnings(user_id), 2):
        ws.cell(row=row, column=1, value=e.get("project_name", ""))
        ws.cell(row=row, column=2, value=e.get("client_name", ""))
        ws.cell(row=row, column=3, value=e.get("amount", 0))
        ws.cell(row=row, column=4, value=e.get("date", ""))
        ws.cell(row=row, column=5, value=e.get("category", ""))
        ws.cell(row=row, column=6, value=e.get("notes", ""))

    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 18

    # ── Projects sheet ──
    ws2 = wb.create_sheet("Projects")
    p_headers = ["Title", "Client", "Budget ($)", "Deadline", "Status", "Priority"]
    for col, h in enumerate(p_headers, 1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
    for row, p in enumerate(get_projects(user_id), 2):
        ws2.cell(row=row, column=1, value=p.get("title", ""))
        ws2.cell(row=row, column=2, value=p.get("client", ""))
        ws2.cell(row=row, column=3, value=p.get("budget", 0))
        ws2.cell(row=row, column=4, value=p.get("deadline", ""))
        ws2.cell(row=row, column=5, value=p.get("status", ""))
        ws2.cell(row=row, column=6, value=p.get("priority", ""))

    # ── Summary sheet ──
    ws3 = wb.create_sheet("Summary")
    stats = get_earnings_stats(user_id)
    t_stats = get_time_stats(user_id)
    summary_rows = [
        ("Total Earnings", f"${stats['total']:,.2f}"),
        ("Monthly Earnings", f"${stats['monthly']:,.2f}"),
        ("Weekly Earnings", f"${stats['weekly']:,.2f}"),
        ("Daily Earnings", f"${stats['daily']:,.2f}"),
        ("Growth Rate", f"{stats['growth_rate']:.1f}%"),
        ("Total Hours Worked", f"{t_stats['total_hours']:.1f}h"),
        ("This Month Hours", f"{t_stats['month_hours']:.1f}h"),
        ("Productivity Score", f"{t_stats['productivity_score']}%"),
    ]
    for i, (k, v) in enumerate(summary_rows, 1):
        ws3.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws3.cell(row=i, column=2, value=v)

    path = os.path.join(_get_report_dir(), f"freelancer_hub_report_{date.today()}.xlsx")
    wb.save(path)
    return path


def export_pdf(user_id: int, report_type: str = "monthly") -> str:
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.units import inch
    except ImportError:
        return export_csv(user_id, report_type)

    path = os.path.join(_get_report_dir(), f"freelancer_hub_{report_type}_{date.today()}.pdf")
    doc = SimpleDocTemplate(path, pagesize=letter)
    styles = getSampleStyleSheet()
    elements = []

    # Title
    title_style = styles["Title"]
    elements.append(Paragraph(f"Freelancer Hub — {report_type.title()} Report", title_style))
    elements.append(Paragraph(f"Generated: {date.today()}", styles["Normal"]))
    elements.append(Spacer(1, 0.3 * inch))

    # Summary
    stats = get_earnings_stats(user_id)
    t_stats = get_time_stats(user_id)
    summary_data = [
        ["Metric", "Value"],
        ["Total Earnings", f"${stats['total']:,.2f}"],
        ["Monthly Earnings", f"${stats['monthly']:,.2f}"],
        ["Weekly Earnings", f"${stats['weekly']:,.2f}"],
        ["Growth Rate", f"{stats['growth_rate']:.1f}%"],
        ["Total Hours", f"{t_stats['total_hours']:.1f}h"],
        ["Productivity Score", f"{t_stats['productivity_score']}%"],
    ]
    t = Table(summary_data, colWidths=[3 * inch, 3 * inch])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1d4ed8")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#f8fafc"), colors.white]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ("PADDING", (0, 0), (-1, -1), 8),
    ]))
    elements.append(Paragraph("Financial Summary", styles["Heading2"]))
    elements.append(t)
    elements.append(Spacer(1, 0.3 * inch))

    # Earnings table
    earnings = get_earnings(user_id)[:20]
    if earnings:
        elements.append(Paragraph("Recent Earnings", styles["Heading2"]))
        e_data = [["Project", "Client", "Amount", "Date", "Category"]]
        for e in earnings:
            e_data.append([
                str(e.get("project_name", ""))[:25],
                str(e.get("client_name", ""))[:20],
                f"${e.get('amount', 0):,.2f}",
                str(e.get("date", "")),
                str(e.get("category", ""))
            ])
        et = Table(e_data, colWidths=[2*inch, 1.5*inch, 1*inch, 1*inch, 1*inch])
        et.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1d4ed8")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#f8fafc"), colors.white]),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("PADDING", (0, 0), (-1, -1), 5),
        ]))
        elements.append(et)

    doc.build(elements)
    return path